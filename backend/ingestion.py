
from pymongo import MongoClient
from config import MONGO_URI, DB_NAME
from sentence_transformers import SentenceTransformer
import fitz, docx, openpyxl, os
from email.parser import BytesParser
from email import policy
# replace the top imports in ingestion.py with this:
try:
    import fitz  # PyMuPDF
except Exception:
    import pymupdf as fitz  # fallback alias

import docx, openpyxl, os

try: import extract_msg; HAS_EXTRACT_MSG=True
except Exception: HAS_EXTRACT_MSG=False
try: from pptx import Presentation; HAS_PPTX=True
except Exception: HAS_PPTX=False

client = MongoClient(MONGO_URI); db = client[DB_NAME]; docs_col = db["docs"]
embedder = SentenceTransformer("all-MiniLM-L6-v2")

def extract_text_from_pdf(path):
    text=""
    try:
        with fitz.open(path) as pdf:
            for p in pdf: text += p.get_text()
    except Exception as e: print("[pdf]", e)
    return text

def extract_text_from_docx(path):
    try:
        d = docx.Document(path); return "\n".join(p.text for p in d.paragraphs)
    except Exception as e: print("[docx]", e); return ""

def extract_text_from_xlsx(path):
    try:
        wb = openpyxl.load_workbook(path, data_only=True); out=[]
        for name in wb.sheetnames:
            ws = wb[name]
            for row in ws.iter_rows(values_only=True):
                out.append(" ".join(str(c) for c in row if c is not None))
        return "\n".join(out)
    except Exception as e: print("[xlsx]", e); return ""

def extract_text_from_pptx(path):
    if not HAS_PPTX: return ""
    try:
        prs = Presentation(path); out=[]
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape,"text"): out.append(shape.text)
        return "\n".join(out)
    except Exception as e: print("[pptx]", e); return ""

def extract_text_from_msg(path):
    if not HAS_EXTRACT_MSG: return ""
    try:
        msg = extract_msg.Message(path)
        body = msg.body or ""; subj = msg.subject or ""
        out = [f"Subject: {subj}", body]
        for att in msg.attachments:
            try:
                name = att.longFilename or att.shortFilename or "attachment"
                data = att.data; ext = os.path.splitext(name)[1].lower()
                import tempfile
                tf = tempfile.NamedTemporaryFile(suffix=ext, delete=False)
                tf.write(data); tf.flush(); tf.close()
                if ext==".pdf": out += [f"[Attachment: {name}]", extract_text_from_pdf(tf.name)]
                elif ext==".docx": out += [f"[Attachment: {name}]", extract_text_from_docx(tf.name)]
                else:
                    try: out += [f"[Attachment: {name}]", data.decode("utf-8","ignore")]
                    except: out += [f"[Attachment: {name}] (binary)"]
                try: os.unlink(tf.name)
                except: pass
            except Exception as e: print("[msg att]", e)
        return "\n".join(out)
    except Exception as e: print("[msg]", e); return ""

def extract_text_from_eml(path):
    try:
        with open(path,"rb") as f:
            m = BytesParser(policy=policy.default).parse(f)
    except Exception as e: print("[eml]", e); return ""
    out=[f"Subject: {m.get('subject','')}", f"From: {m.get('from','')}", f"To: {m.get('to','')}",""]
    body=""
    if m.is_multipart():
        for part in m.walk():
            ctype = part.get_content_type()
            disp = part.get_content_disposition()
            if ctype=="text/plain" and disp in (None,"inline"):
                try: body += part.get_content()
                except: body += part.get_payload(decode=True).decode("utf-8","ignore")
            elif disp=="attachment":
                name = part.get_filename() or "attachment"
                ext = os.path.splitext(name)[1].lower()
                data = part.get_payload(decode=True) or b""
                import tempfile
                tf = tempfile.NamedTemporaryFile(suffix=ext, delete=False)
                tf.write(data); tf.flush(); tf.close()
                if ext==".pdf": out += [f"[Attachment: {name}]", extract_text_from_pdf(tf.name)]
                elif ext==".docx": out += [f"[Attachment: {name}]", extract_text_from_docx(tf.name)]
                else:
                    try: out += [f"[Attachment: {name}]", data.decode("utf-8","ignore")]
                    except: out += [f"[Attachment: {name}] (binary)"]
                try: os.unlink(tf.name)
                except: pass
            elif ctype=="text/html" and not body.strip():
                try:
                    html = part.get_content()
                    import re
                    body = re.sub("<[^<]+?>","", html)
                except: pass
    else:
        try: body = m.get_content()
        except: body = m.get_payload(decode=True).decode("utf-8","ignore")
    out.append(body); return "\n".join(out)

def extract_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext==".pdf": return extract_text_from_pdf(path)
    if ext==".docx": return extract_text_from_docx(path)
    if ext in [".xlsx",".xls"]: return extract_text_from_xlsx(path)
    if ext==".pptx": return extract_text_from_pptx(path)
    if ext==".msg": return extract_text_from_msg(path)
    if ext==".eml": return extract_text_from_eml(path)
    try:
        with open(path,"r",encoding="utf-8",errors="ignore") as f: return f.read()
    except: return ""

def chunk_text(text, chunk_size=800, overlap=100):
    words = text.split()
    if not words: return []
    chunks=[]; i=0
    while i < len(words):
        chunks.append(" ".join(words[i:i+chunk_size]))
        i += (chunk_size - overlap)
    return chunks

def ingest_file(path: str):
    text = extract_text(path)
    if not text.strip(): return 0
    chunks = chunk_text(text)
    if not chunks: return 0
    embeddings = embedder.encode(chunks, convert_to_numpy=True).tolist()
    for ch, emb in zip(chunks, embeddings):
        docs_col.insert_one({"filename": os.path.basename(path), "content": ch, "embedding": emb})
    return len(chunks)
